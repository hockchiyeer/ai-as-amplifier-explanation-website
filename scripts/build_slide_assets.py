#!/usr/bin/env python3
"""
Rebuild the self-contained slide "screenshots" used by the tutorial.

For every deck listed in DECKS this script
  1. copies resources/data/pptx_images/<folder>/SlideN.PNG  ->  client/public/slides/<id>/slide-NN.png
  2. extracts each slide's title and text (python-pptx) into client/src/lib/slides-data.json,
     which powers image alt text and the "Text version of this slide" panel.

Run it again whenever a deck or its slide images change:
    pip install python-pptx
    python3 scripts/build_slide_assets.py
"""
import glob, json, os, shutil
from pptx import Presentation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
# id (must match pptShowcase.items[].id in i18n) -> pptx file stem (also the image folder name)
DECKS = {
    "amplifier": "AI_As_Amplifier_Multiplier_Compounder_Chinese_version_35_slides",
    "testing": "AI_in_Software_Test_Automation_ByErHockChiye",
    "politics": "AI政治策略课程_完整版",
}

def slide_text(slide):
    title, lines = None, []
    if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
        title = slide.shapes.title.text_frame.text.strip()
    for sh in slide.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                t = "".join(r.text for r in p.runs).strip()
                if t:
                    lines.append(t)
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                lines.append(" | ".join(c.text.strip() for c in row.cells))
    return (title or (lines[0] if lines else "")).replace("\n", " ").strip(), lines

out = {}
for deck_id, stem in DECKS.items():
    pptx_path = os.path.join(ROOT, "resources/data/pptx", stem + ".pptx")
    img_dir = os.path.join(ROOT, "resources/data/pptx_images", stem)
    prs = Presentation(pptx_path)
    count = len(glob.glob(os.path.join(img_dir, "Slide*.PNG")))
    if count != len(prs.slides):
        raise SystemExit(f"{deck_id}: {len(prs.slides)} slides but {count} images in {img_dir}")
    dst = os.path.join(ROOT, "client/public/slides", deck_id)
    os.makedirs(dst, exist_ok=True)
    for i in range(1, count + 1):
        shutil.copyfile(os.path.join(img_dir, f"Slide{i}.PNG"), os.path.join(dst, f"slide-{i:02d}.png"))
    slides = []
    for i, s in enumerate(prs.slides, 1):
        title, lines = slide_text(s)
        slides.append({"n": i, "title": title, "text": lines})
    out[deck_id] = {"file": stem + ".pptx", "count": count, "slides": slides}
    print(f"{deck_id}: {count} slides")

with open(os.path.join(ROOT, "client/src/lib/slides-data.json"), "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, separators=(",", ":"))
